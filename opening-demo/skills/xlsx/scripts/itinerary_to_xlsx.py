"""Build cypher_plan.xlsx from itinerary.json (written by submit_itinerary)."""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="161B36")
HEADER_FONT = Font(bold=True, color="FFFFFF")
BAND_FILL = PatternFill("solid", fgColor="F2F3FA")
HALL_FILLS = {1: "FDE9C2", 2: "D4F4EF", 3: "E3E0F7"}
THIN = Side(style="thin", color="D0D3E6")
WRAP = Alignment(wrap_text=True, vertical="top")
DAY_COLUMNS = [("Time", 13), ("Hall", 7), ("Category", 16), ("Session id", 13), ("Title", 46),
               ("Speakers", 30), ("Why it fits", 46), ("Question to ask", 46)]


def style_header(ws, widths: list[tuple[str, int]]) -> None:
    for col, (name, width) in enumerate(widths, start=1):
        cell = ws.cell(row=1, column=col, value=name)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(vertical="center")
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 22


def main() -> int:
    p = argparse.ArgumentParser()
    p.add_argument("--itinerary", default="itinerary.json")
    p.add_argument("--out", default="cypher_plan.xlsx")
    p.add_argument("--spotlight", default=None, help="JSON list of researched speakers")
    a = p.parse_args()

    path = Path(a.itinerary)
    if not path.exists():
        print(f"ERROR: {path} not found. Call submit_itinerary first; it writes itinerary.json.")
        return 1
    itin = json.loads(path.read_text())

    wb = Workbook()
    overview = wb.active
    overview.title = "Overview"

    day_sheets = []
    for day in sorted(itin.get("days", []), key=lambda d: d.get("day", 0)):
        ws = wb.create_sheet(f"Day {day['day']}")
        day_sheets.append(ws.title)
        style_header(ws, DAY_COLUMNS)
        items = sorted(day.get("items", []), key=lambda i: i.get("start", ""))
        for r, item in enumerate(items, start=2):
            speakers = ", ".join(s.get("name", "") for s in item.get("speakers", []))
            values = [f"{item.get('start', '')}–{item.get('end', '')}", item.get("hall"), item.get("category", ""),
                      item.get("session_id", ""), item.get("title", ""), speakers,
                      item.get("reason", ""), item.get("question", "")]
            for c, v in enumerate(values, start=1):
                cell = ws.cell(row=r, column=c, value=v)
                cell.alignment, cell.border = WRAP, Border(bottom=THIN)
                if r % 2 == 0:
                    cell.fill = BAND_FILL
            hall = item.get("hall")
            if hall in HALL_FILLS:
                ws.cell(row=r, column=2).fill = PatternFill("solid", fgColor=HALL_FILLS[hall])
        print(f"{ws.title}: {len(items)} items")

    # Overview with live formulas so totals survive edits.
    style_header(overview, [("Day", 10), ("Items", 10), ("Workshops", 12), ("Panels & debates", 18), ("Lunch", 10)])
    for r, name in enumerate(day_sheets, start=2):
        rng = f"'{name}'!C2:C200"
        overview.cell(row=r, column=1, value=name)
        overview.cell(row=r, column=2, value=f"=COUNTA('{name}'!D2:D200)")
        overview.cell(row=r, column=3, value=f'=COUNTIF({rng},"*Workshop*")')
        overview.cell(row=r, column=4, value=f'=COUNTIF({rng},"Panel*")+COUNTIF({rng},"Debate")')
        overview.cell(row=r, column=5, value=f'=COUNTIF({rng},"Break / Lunch")')
    total = len(day_sheets) + 2
    overview.cell(row=total, column=1, value="Total").font = Font(bold=True)
    for c in range(2, 6):
        col = get_column_letter(c)
        overview.cell(row=total, column=c, value=f"=SUM({col}2:{col}{total - 1})").font = Font(bold=True)
    if itin.get("summary"):
        overview.cell(row=total + 2, column=1, value="Summary").font = Font(bold=True)
        overview.cell(row=total + 3, column=1, value=itin["summary"]).alignment = WRAP
        overview.merge_cells(start_row=total + 3, start_column=1, end_row=total + 3, end_column=5)
        overview.row_dimensions[total + 3].height = 90

    if a.spotlight:
        try:
            spotlight = json.loads(a.spotlight)
        except json.JSONDecodeError as exc:
            print(f"ERROR: --spotlight is not valid JSON: {exc}")
            return 1
        ws = wb.create_sheet("Speakers")
        style_header(ws, [("Speaker", 26), ("Company", 24), ("Headline", 50), ("Why meet them", 50)])
        for r, s in enumerate(spotlight, start=2):
            for c, key in enumerate(["name", "company", "headline", "why_meet"], start=1):
                ws.cell(row=r, column=c, value=s.get(key, "")).alignment = WRAP
        print(f"Speakers: {len(spotlight)} rows")

    wb.save(a.out)
    print(f"Saved {a.out} with sheets: {', '.join(wb.sheetnames)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
