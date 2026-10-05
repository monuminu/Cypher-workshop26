"""Offline integration tests with the real pinned framework and scripted model responses."""
import asyncio
import json
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import httpx
import app as demo
from interpreter import PythonInterpreter
from agent_framework import (BaseChatClient, ChatMiddlewareLayer, FunctionInvocationLayer,
                             ChatResponse, ChatResponseUpdate, ResponseStream, Content, Message, SkillsProvider)


class ScriptedClient(FunctionInvocationLayer, ChatMiddlewareLayer, BaseChatClient):
    def __init__(self, steps=(), gate=None):
        super().__init__()
        self.steps, self.gate, self.calls = list(steps), gate, 0
        self.seen_tools = set()

    def _inner_get_response(self, *, messages, stream, options, **kwargs):
        assert stream, "The leaf provider must receive real streaming requests."
        calls, results = [], []
        for message in messages:
            for content in message.contents:
                if content.type == 'function_call': calls.append(content.call_id)
                elif content.type == 'function_result': results.append(content.call_id)
        assert len(calls) == len(set(calls)), 'Duplicated tool call in model history'
        assert len(results) == len(set(results)), 'Duplicated tool output in model history'
        assert set(calls) == set(results), 'Unpaired tool call/output in model history'
        self.calls += 1
        self.seen_tools.update(getattr(t, "name", "") for t in options.get("tools", []))
        step = self.steps.pop(0) if self.steps else None
        call_id = f"call-{self.calls}"
        async def updates():
            if self.gate and self.calls == 1:
                await self.gate()
            if step:
                args = json.dumps(step[1])
                split = max(1, len(args)//2)
                for chunk in (args[:split], args[split:]):
                    yield ChatResponseUpdate(role='assistant', contents=[Content.from_function_call(call_id, step[0], arguments=chunk)])
            else:
                for text in ('Finished ', 'scripted ', 'test.'):
                    yield ChatResponseUpdate(role='assistant', contents=[Content.from_text(text)])
            await asyncio.sleep(0)
        return ResponseStream(updates(), finalizer=ChatResponse.from_updates)


CREATE_FILES = '''from pathlib import Path
from openpyxl import Workbook, load_workbook
from pptx import Presentation
assert Path('inputs/source.txt').read_text() == 'same source'
w=Workbook(); w.active['A1']='User-generated example'; w.save('outputs/example.xlsx')
assert load_workbook('outputs/example.xlsx').active['A1'].value == 'User-generated example'
p=Presentation(); p.slides.add_slide(p.slide_layouts[0]); p.save('outputs/example.pptx')
assert len(Presentation('outputs/example.pptx').slides)==1
print('Both files reopened')
'''


class AppTests(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        self.root_patch = patch.object(demo, "RUNS", self.root)
        self.root_patch.start()
        demo.live.clear()
        self.http = httpx.AsyncClient(transport=httpx.ASGITransport(app=demo.app),
                                     base_url="http://testserver", headers={"X-Opening-Demo": "1"})

    async def asyncTearDown(self):
        await asyncio.gather(*(r.close() for r in demo.live.values()), return_exceptions=True)
        demo.live.clear()
        await self.http.aclose()
        self.root_patch.stop()
        self.tmp.cleanup()

    async def create(self, clients, **data):
        with patch.object(demo, "client_factory", side_effect=clients):
            response = await self.http.post('/api/runs', data={"prompt": "Create files from my inputs.", **data},
                files={"files": ("source.txt", b"same source")})
        self.assertEqual(response.status_code, 200, response.text)
        return response.json()['id']

    async def finish(self, ident):
        await asyncio.wait_for(asyncio.gather(*(s.task for s in demo.live[ident].sides.values())), 20)
        return (await self.http.get(f'/api/runs/{ident}')).json()

    async def test_parallel_generic_artifacts_skills_and_downloads(self):
        arrivals = 0
        ready = asyncio.Event()
        async def gate():
            nonlocal arrivals
            arrivals += 1
            if arrivals == 2:
                ready.set()
            await ready.wait()
        basic = ScriptedClient([('python_execute', {'code': CREATE_FILES}), None], gate)
        harness = ScriptedClient([('load_skill', {'skill_name': 'xlsx'}),
                                  ('load_skill', {'skill_name': 'pptx'}),
                                  ('python_execute', {'code': CREATE_FILES}), None], gate)
        ident = await self.create([basic, harness])
        data = await self.finish(ident)
        self.assertEqual(arrivals, 2)
        self.assertEqual(basic.seen_tools, {'python_execute'})
        self.assertIn('load_skill', harness.seen_tools)
        self.assertEqual(data['sides']['harness']['skills'], ['pptx', 'xlsx'])
        for name in ('basic', 'harness'):
            self.assertEqual(data['sides'][name]['status'], 'finished', data['sides'][name])
            self.assertEqual({f['name'] for f in data['sides'][name]['files']}, {'example.pptx', 'example.xlsx'})
            download = await self.http.get(f'/api/runs/{ident}/files/{name}/example.pptx')
            self.assertEqual(download.status_code, 200)
            self.assertTrue(download.content.startswith(b'PK'))
        self.assertTrue(any(e['kind']=='tool_end' for e in data['events']))
        self.assertEqual((await self.http.get(f'/api/runs/{ident}?after={data["cursor"]}')).json()['events'], [])
        denied = await self.http.get(f'/api/runs/{ident}/files/basic/%2e%2e%2finputs%2fsource.txt')
        self.assertEqual(denied.status_code, 404)

    async def test_failure_does_not_stop_other_agent_and_call_limit_visible(self):
        async def fail():
            raise RuntimeError('Test provider failure')
        ident = await self.create([ScriptedClient(gate=fail), ScriptedClient([
            ('python_execute', {'code': "print('hello')"}), None])], max_calls=1)
        data = await self.finish(ident)
        self.assertEqual(data['sides']['basic']['status'], 'error')
        self.assertEqual(data['sides']['harness']['status'], 'call limit')

    async def test_approval_can_be_declined(self):
        ident = await self.create([ScriptedClient(), ScriptedClient([
            ('python_execute', {'code': "raise RuntimeError('must not run')"}), None])], ask_tools='true')
        harness = demo.live[ident].sides['harness']
        async def wait_pending():
            while not harness.pending:
                await asyncio.sleep(.01)
        await asyncio.wait_for(wait_pending(), 5)
        ident_approval = next(iter(harness.pending))
        response = await self.http.post(f'/api/runs/{ident}/approve/{ident_approval}', json={'allow':False})
        self.assertEqual(response.status_code, 200)
        data = await self.finish(ident)
        self.assertEqual(data['sides']['harness']['status'], 'finished')
        self.assertFalse(list((demo.live[ident].directory/'logs/harness').glob('*.py')))

    async def test_stop_and_resume_persisted_history(self):
        wait = asyncio.Event()
        async def block(): await wait.wait()
        ident = await self.create([ScriptedClient(gate=block), ScriptedClient(gate=block)])
        await asyncio.sleep(.01)
        response = await self.http.post(f'/api/runs/{ident}/stop')
        self.assertEqual(response.status_code, 200)
        data = await self.finish(ident)
        self.assertEqual([s['status'] for s in data['sides'].values()], ['stopped','stopped'])
        await demo.live.pop(ident).close()
        self.assertFalse((await self.http.get(f'/api/runs/{ident}')).json()['busy'])
        with patch.object(demo, 'client_factory', side_effect=[ScriptedClient(),ScriptedClient()]):
            response=await self.http.post(f'/api/runs/{ident}/continue',json={'prompt':'Continue the task.'})
        self.assertEqual(response.status_code,200,response.text)
        data=await self.finish(ident)
        self.assertTrue(all(s['status']=='finished' for s in data['sides'].values()), data)

    async def test_cross_origin_and_missing_header_rejected(self):
        response = await self.http.post('/api/runs', headers={'Origin':'https://example.com'},data={'prompt':'x'})
        self.assertEqual(response.status_code,403)
        async with httpx.AsyncClient(transport=httpx.ASGITransport(app=demo.app),base_url='http://testserver') as client:
            self.assertEqual((await client.post('/api/runs',data={'prompt':'x'})).status_code,403)

    async def test_harness_todo_continuation_and_followup_keep_history(self):
        harness = ScriptedClient([
            ('todos_add', {'todos':[{'title':'Review the input', 'description':'Offline test'}]}), None,
            ('todos_complete', {'items':[{'id':1,'reason':'Input reviewed'}]}), None])
        ident = await self.create([ScriptedClient(), harness])
        data = await self.finish(ident)
        self.assertEqual(data['sides']['harness']['model_calls'],4)
        self.assertEqual(data['sides']['harness']['todos'],[{'title':'Review the input','complete':True}])
        self.assertEqual(data['sides']['harness']['status'],'finished')
        sessions={name:side.session.session_id for name,side in demo.live[ident].sides.items()}
        response=await self.http.post(f'/api/runs/{ident}/continue',json={'prompt':'Explain what you did.'})
        self.assertEqual(response.status_code,200)
        await self.finish(ident)
        self.assertEqual(sessions,{name:side.session.session_id for name,side in demo.live[ident].sides.items()})

    async def test_runtime_has_no_agenda_dependency_or_exporter(self):
        import ast
        root=Path(__file__).resolve().parents[1]
        for name in ('runtime.py','interpreter.py','app.py'):
            tree=ast.parse((root/name).read_text())
            imports=[node.module for node in ast.walk(tree) if isinstance(node,ast.ImportFrom)]
            self.assertFalse(any('agenda' in (name or '') for name in imports))
        ident=await self.create([ScriptedClient(),ScriptedClient()])
        await self.finish(ident)
        for side in demo.live[ident].sides.values():
            self.assertNotIn('agenda',side.interpreter.python_execute.__doc__.lower())
            self.assertNotIn('workbook',side.instructions.lower())

    async def test_selected_model_applies_to_both_and_survives_restore(self):
        with patch.object(demo,'config_info',return_value={'provider':'openai','model':'configured-model'}):
            config=(await self.http.get('/api/config')).json()
            self.assertIn('gpt-4o',[m['value'] for m in config['models']])
            with patch.object(demo,'client_factory',side_effect=[ScriptedClient(),ScriptedClient()]) as factory:
                response=await self.http.post('/api/runs',data={'prompt':'Explain agents.','model':'gpt-4o'})
            self.assertEqual(response.status_code,200,response.text)
            self.assertEqual([call.kwargs for call in factory.call_args_list],[{'model':'gpt-4o'}]*2)
        ident=response.json()['id']
        data=await self.finish(ident)
        self.assertEqual(data['config']['model'],'gpt-4o')
        await demo.live.pop(ident).close()
        with patch.object(demo,'config_info',return_value={'provider':'openai','model':'new-default'}):
            with patch.object(demo,'client_factory',side_effect=[ScriptedClient(),ScriptedClient()]) as factory:
                response=await self.http.post(f'/api/runs/{ident}/continue',json={'prompt':'Continue.'})
            self.assertEqual(response.status_code,200,response.text)
            self.assertEqual([call.kwargs for call in factory.call_args_list],[{'model':'gpt-4o'}]*2)
        await self.finish(ident)

    async def test_unsupported_model_rejected_without_client_creation(self):
        with patch.object(demo,'client_factory') as factory:
            response=await self.http.post('/api/runs',data={'prompt':'Explain agents.','model':'not-offered'})
            self.assertEqual(response.status_code,422)
            factory.assert_not_called()

    async def test_client_factory_passes_model_override_to_workshop_client(self):
        with patch('workshop_utils.get_chat_client') as client:
            demo.client_factory(model='gpt-4o')
            client.assert_called_once_with(model='gpt-4o')

    async def test_deltas_arrive_before_completion_and_sse_replays_once(self):
        release = asyncio.Event()
        started = asyncio.Event()
        class SlowTextClient(ScriptedClient):
            def _inner_get_response(self, *, messages, stream, options, **kwargs):
                assert stream
                async def updates():
                    yield ChatResponseUpdate(role='assistant',contents=[Content.from_text('First chunk. ')])
                    started.set()
                    await release.wait()
                    yield ChatResponseUpdate(role='assistant',contents=[Content.from_text('Last chunk.')])
                return ResponseStream(updates(),finalizer=ChatResponse.from_updates)
        ident = await self.create([SlowTextClient(),SlowTextClient()])
        await asyncio.wait_for(started.wait(),5)
        run = demo.live[ident]
        self.assertTrue(run.busy)
        self.assertTrue(any(e['kind']=='text_delta' for e in run.events))
        self.assertFalse(any(e['kind']=='answer' for e in run.events))
        partial=(await self.http.get(f'/api/runs/{ident}')).json()
        for side in partial['sides'].values():
            self.assertEqual(side['answer'],'First chunk. ')
            self.assertNotIn('seconds',side)
        release.set()
        data=await self.finish(ident)
        self.assertEqual(data['sides']['basic']['answer'],'First chunk. Last chunk.')
        response=await self.http.get(f'/api/runs/{ident}/stream')
        self.assertIn('text/event-stream',response.headers['content-type'])
        frames=response.text.split('\n\n')
        ids=[int(frame.splitlines()[0].split(': ')[1]) for frame in frames if frame.startswith('id:')]
        self.assertEqual(ids,list(range(len(run.events))))
        self.assertIn('event: idle',response.text)
        last=ids[-2]
        replay=await self.http.get(f'/api/runs/{ident}/stream',headers={'Last-Event-ID':str(last)})
        self.assertEqual(replay.text.count('event: activity'),1)

    async def test_stop_mid_stream_keeps_visible_partial_text(self):
        started = asyncio.Event()
        class EndlessClient(ScriptedClient):
            def _inner_get_response(self, **kwargs):
                async def updates():
                    yield ChatResponseUpdate(role='assistant',contents=[Content.from_text('Partial response')])
                    started.set()
                    await asyncio.Event().wait()
                return ResponseStream(updates(),finalizer=ChatResponse.from_updates)
        ident=await self.create([EndlessClient(),EndlessClient()])
        await asyncio.wait_for(started.wait(),5)
        await self.http.post(f'/api/runs/{ident}/stop')
        data=await self.finish(ident)
        for side in data['sides'].values():
            self.assertEqual(side['status'],'stopped')
            self.assertEqual(side['answer'],'Partial response')

    async def test_python_cancellation_stops_process_and_has_no_deadline(self):
        runner=PythonInterpreter(self.root/'workspace',self.root/'logs')
        task=asyncio.create_task(runner.execute("import time; from pathlib import Path; Path('started').write_text('yes'); time.sleep(60); Path('late').write_text('bad')"))
        async def wait_started():
            while not (self.root/'workspace/started').exists(): await asyncio.sleep(.01)
        await asyncio.wait_for(wait_started(),5)
        task.cancel()
        with self.assertRaises(asyncio.CancelledError): await task
        self.assertFalse((self.root/'workspace/late').exists())


if __name__=='__main__': unittest.main()
