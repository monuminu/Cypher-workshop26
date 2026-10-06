import json
import shutil
import subprocess
import sys
from pathlib import Path

import pytest

from demo.data import personas
from demo.tools import DECK_FILE, OUT_DIR, XLSX_FILE, RunState, build_tools

SKILLS = Path(__file__).resolve().parent.parent / "skills"

# A small, valid Day-1 plan for Priya (real session ids).
ITINERARY = {
    "days": [{"day": 1, "items": [
        {"session_id": "D1-H3-1100", "title": "Your Agent Stack Is Broken at the Bottom",
         "reason": "Agent reliability from the ground up.", "question": "What breaks first?"},
        {"session_id": "D1-H3-1230", "title": "Lunch"},
        {"session_id": "D1-H3-1930", "title": "Agent Harness for Enterprise",
         "reason": "Her must-attend.", "question": "Which harness feature pays off first?"},
    ]}],
    "summary": "Agents first.",
}


@pytest.fixture
def tools():
    state = RunState(persona=personas()["engineer"], side="baseline", run_id="test")
    yield state, {t.name: t.func for t in build_tools(state)}
    shutil.rmtree(state.workspace, ignore_errors=True)


def test_list_sessions_filters(tools):
    _, t = tools
    rows = t["list_sessions"](day=1, hall=3).splitlines()
    assert rows and all(r.startswith("D1-H3-") for r in rows)
    assert "Agent Harness for Enterprise" in t["list_sessions"](day=1, category="workshop")


def test_search_finds_harness_workshop(tools):
    assert "D1-H3-1930" in tools[1]["search_sessions"](query="agent harness")


def test_get_session_records_reads(tools):
    state, t = tools
    assert "Manoranjan" in t["get_session"](session_id="D1-H3-1930")
    assert "D1-H3-1930" in state.fetched_ids
    assert "No session" in t["get_session"](session_id="D9-H9-0000")


def test_lookup_speaker_records_research(tools):
    state, t = tools
    out = t["lookup_speaker"](name="Manoranjan")
    assert "Bio:" in out and "D1-H3-1930" in out
    assert "Manoranjan Rajguru" in state.researched


def test_run_python_runs_in_workspace(tools):
    state, t = tools
    out = t["run_python"](code="open('x.txt','w').write('hi'); print('ok')")
    assert "ok" in out and "x.txt" in out


def test_submit_writes_calendar_and_enriched_itinerary(tools):
    state, t = tools
    t["submit_itinerary"](itinerary_json=json.dumps(ITINERARY))
    assert (OUT_DIR / state.ics_path).read_text().count("BEGIN:VEVENT") == 3
    enriched = json.loads((state.workspace / "itinerary.json").read_text())
    assert all("start" in i and "hall" in i for d in enriched["days"] for i in d["items"])


def test_skill_scripts_build_deliverables(tools):
    state, t = tools
    t["submit_itinerary"](itinerary_json=json.dumps(ITINERARY))
    spot = json.dumps([{"name": "Manoranjan Rajguru", "company": "Microsoft", "headline": "x", "why_meet": "y"}])
    for script, extra in [("xlsx/scripts/itinerary_to_xlsx.py", []),
                          ("pptx/scripts/build_briefing_deck.py", ["--attendee", "Priya"])]:
        subprocess.run([sys.executable, str(SKILLS / script), "--spotlight", spot, *extra],
                       cwd=state.workspace, check=True, capture_output=True)
    from openpyxl import load_workbook
    from pptx import Presentation

    wb = load_workbook(state.workspace / XLSX_FILE)
    assert {"Overview", "Day 1", "Speakers"} <= set(wb.sheetnames)
    assert len(Presentation(str(state.workspace / DECK_FILE)).slides) == 4  # title, day 1, spotlight, questions


def test_profile_mentions_commitments_and_deliverables(tools):
    profile = tools[1]["get_attendee_profile"]()
    assert "Booth duty" in profile and "D1-H3-1930" in profile and XLSX_FILE in profile
