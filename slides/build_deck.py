"""Build the Cypher 2026 Agentic AI companion lecture deck.

Run: python slides/build_deck.py
Requires: python-pptx and Pillow (install from slides/requirements.txt).
Uses the workshop diagrams in docs/assets; no external template is needed.
"""

from __future__ import annotations

import os

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
ASSETS = os.path.join(ROOT, "docs", "assets")
OUT = os.path.join(HERE, "Cypher-2026-Agentic-AI.pptx")

# --- brand palette --------------------------------------------------------
INK = RGBColor(0x18, 0x18, 0x18)       # near-black body text on white
GRAY = RGBColor(0x44, 0x44, 0x44)      # secondary text
INDIGO = RGBColor(0x4B, 0x3F, 0xC4)    # matches the diagrams' indigo
MAGENTA = RGBColor(0xD6, 0x24, 0x9F)   # brand magenta
CYAN = RGBColor(0x10, 0x9A, 0xB5)      # brand cyan (darkened for white-bg legibility)
PURPLE = RGBColor(0x6B, 0x3F, 0xA0)    # brand purple
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CARD_BG = RGBColor(0xF3, 0xF1, 0xFB)   # very light purple card fill
FONT = "Arial"

# slide canvas (inches)
SW, SH = 10.0, 5.625
DIAG_RATIO = 1.5  # all diagrams are 3:2

# ============================ low-level helpers ============================
def _set_run(run, text, size, color, bold=False, italic=False):
    run.text = text
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color


def add_textbox(slide, x, y, w, h, *, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    return box


def add_bullets(slide, x, y, w, h, items, *, size=15, color=INK, gap=6,
                bullet_color=INDIGO):
    """items: list of (text) or (text, level). Renders dash bullets."""
    box = add_textbox(slide, x, y, w, h)
    tf = box.text_frame
    first = True
    for it in items:
        text, level = (it if isinstance(it, tuple) else (it, 0))
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.level = level
        p.space_after = Pt(gap)
        # bullet glyph
        rb = p.add_run()
        _set_run(rb, ("•  " if level == 0 else "–  "),
                 size, bullet_color, bold=True)
        rt = p.add_run()
        _set_run(rt, text, size if level == 0 else size - 1,
                 color if level == 0 else GRAY)
    return box


def add_picture_fit(slide, path, *, x, y, max_w, max_h):
    """Place an image scaled to fit inside (max_w, max_h), centered in that box."""
    iw, ih = Image.open(path).size
    ratio = iw / ih
    w = max_w
    h = w / ratio
    if h > max_h:
        h = max_h
        w = h * ratio
    cx = x + (max_w - w) / 2
    cy = y + (max_h - h) / 2
    return slide.shapes.add_picture(path, Inches(cx), Inches(cy),
                                    Inches(w), Inches(h))


def add_caption(slide, x, y, w, text, *, size=11, color=GRAY, align=PP_ALIGN.CENTER):
    box = add_textbox(slide, x, y, w, 0.3)
    p = box.text_frame.paragraphs[0]
    p.alignment = align
    _set_run(p.add_run(), text, size, color, italic=True)
    return box


def add_rounded(slide, x, y, w, h, fill):
    from pptx.enum.shapes import MSO_SHAPE
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                 Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def add_pill(slide, x, y, w, h, text, fill, *, tcolor=WHITE, size=12):
    shp = add_rounded(slide, x, y, w, h, fill)
    tf = shp.text_frame
    tf.word_wrap = True
    tf.margin_top = Pt(2)
    tf.margin_bottom = Pt(2)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    _set_run(p.add_run(), text, size, tcolor, bold=True)
    return shp


# ============================ slide builders ==============================
def new_slide(prs, *, dark=False):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = INDIGO if dark else WHITE
    return slide


def title_slide(prs):
    s = new_slide(prs, dark=True)
    box = add_textbox(s, 0.6, 0.65, 8.8, 0.4)
    _set_run(box.text_frame.paragraphs[0].add_run(), "Cypher 2026", 18, WHITE, bold=True)
    box = add_textbox(s, 0.6, 1.65, 8.8, 1.65)
    _set_run(box.text_frame.paragraphs[0].add_run(),
             "Agentic AI: Building, Optimizing\n& Operationalizing Agents", 30, WHITE, bold=True)
    box = add_textbox(s, 0.6, 3.9, 8.8, 0.9)
    _set_run(box.text_frame.paragraphs[0].add_run(), "Manoranjan Rajguru", 18, WHITE, bold=True)
    _set_run(box.text_frame.add_paragraph().add_run(), "Microsoft", 15, WHITE)


def section_slide(prs, label, subtitle=None):
    s = new_slide(prs)
    box = add_textbox(s, 0.6, 1.8, 8.8, 0.9)
    _set_run(box.text_frame.paragraphs[0].add_run(), label, 30, INDIGO, bold=True)
    if subtitle:
        box = add_textbox(s, 0.6, 2.85, 8.8, 0.7)
        _set_run(box.text_frame.paragraphs[0].add_run(), subtitle, 18, GRAY)
    return s


def content_slide(prs, title, *, title_color=INDIGO):
    s = new_slide(prs)
    box = add_textbox(s, 0.55, 0.4, 8.9, 0.8)
    _set_run(box.text_frame.paragraphs[0].add_run(), title, 24, title_color, bold=True)
    return s


def text_and_diagram(prs, title, bullets, diagram, caption=None):
    """Left: bullets. Right: diagram."""
    s = content_slide(prs, title)
    add_bullets(s, 0.55, 1.35, 4.55, 3.9, bullets, size=14, gap=7)
    if diagram:
        add_picture_fit(s, os.path.join(ASSETS, diagram),
                        x=5.25, y=1.25, max_w=4.45, max_h=3.5)
        if caption:
            add_caption(s, 5.25, 4.8, 4.45, caption)
    return s


def diagram_hero(prs, title, diagram, bullets_below=None, caption=None):
    """Big centered diagram with optional one-line takeaways underneath."""
    s = content_slide(prs, title)
    add_picture_fit(s, os.path.join(ASSETS, diagram),
                    x=2.3, y=1.15, max_w=5.4, max_h=3.2)
    if caption:
        add_caption(s, 2.3, 4.35, 5.4, caption)
    if bullets_below:
        add_bullets(s, 0.6, 4.6, 8.8, 0.9, bullets_below, size=12, gap=2)
    return s


# ============================ deck assembly ===============================
def build():
    prs = Presentation()
    prs.slide_width = Inches(SW)
    prs.slide_height = Inches(SH)

    # 1) Title
    title_slide(prs)

    # 2) Agenda
    s = content_slide(prs, "The Day — From One LLM Call to a Full Agent")
    rows = [
        ("M1", "Your First Agent", "the agent loop, streaming"),
        ("M2", "Tools & Function Calling", "give the agent hands"),
        ("M3", "Context Engineering", "sessions, memory, compaction"),
        ("M4", "The Agent Harness  ★", "create_harness_agent — batteries included"),
        ("M5", "Multi-Agent Orchestration", "sequential, concurrent, handoff…"),
        ("M6", "Evaluating & Optimizing", "measure quality, gate CI"),
        ("M7", "Operationalizing", "tracing, middleware, guardrails"),
        ("M8", "Capstone & Hosting", "combine it all; A2A, Functions"),
    ]
    y = 1.3
    for code, name, desc in rows:
        add_pill(s, 0.55, y, 0.7, 0.34, code, INDIGO, size=12)
        bx = add_textbox(s, 1.4, y - 0.02, 8.1, 0.4, anchor=MSO_ANCHOR.MIDDLE)
        p = bx.text_frame.paragraphs[0]
        _set_run(p.add_run(), f"{name}  ", 14, INK, bold=True)
        _set_run(p.add_run(), f"— {desc}", 12, GRAY)
        y += 0.43
    add_caption(s, 0.55, y + 0.02, 9.0,
                "Labs are hands-on Jupyter notebooks · live at monuminu.github.io/Cypher-workshop26",
                align=PP_ALIGN.LEFT)

    # 3) Section: Foundations
    section_slide(prs, "Foundations", "What is an agent, and why now?")

    # 4) What is an agent
    text_and_diagram(
        prs, "From an LLM Call to an Agent",
        [
            "An LLM alone is text-in, text-out — no memory, no actions.",
            "An agent wraps the model in a loop that adds three powers:",
            ("Tools — it can act, not just talk", 1),
            ("Memory — it remembers across turns", 1),
            ("Control flow — it plans and repeats until done", 1),
            "Agent = model + instructions + tools + a loop.",
        ],
        "agent-anatomy.png",
        caption="The anatomy of an agent",
    )

    # 5) The agent loop
    s = content_slide(prs, "The Agent Loop (ReAct)")
    add_bullets(s, 0.55, 1.35, 8.9, 1.4, [
        "Every agent.run() executes a reason → act → observe loop:",
        ("1 · Send conversation + available tools to the model", 1),
        ("2 · If the model calls a tool, run it and feed the result back", 1),
        ("3 · Repeat until the model returns a final answer", 1),
    ], size=14, gap=6)
    # simple flow boxes
    flow = [("REASON", INDIGO), ("ACT (tool)", MAGENTA), ("OBSERVE", CYAN), ("ANSWER", PURPLE)]
    fx = 0.7
    for i, (lbl, col) in enumerate(flow):
        add_pill(s, fx, 3.3, 1.85, 0.7, lbl, col, size=13)
        fx += 2.1
        if i < len(flow) - 1:
            ar = add_textbox(s, fx - 0.27, 3.42, 0.3, 0.4)
            _set_run(ar.text_frame.paragraphs[0].add_run(), "→", 20, GRAY, bold=True)
    add_caption(s, 0.55, 4.4, 8.9,
                "With no tools the loop is a single hop; add tools (M2) and it iterates.",
                align=PP_ALIGN.LEFT)

    # 6) Why Microsoft Agent Framework
    s = content_slide(prs, "Why the Microsoft Agent Framework")
    add_bullets(s, 0.55, 1.35, 4.7, 3.9, [
        "Open-source SDK — the successor to Semantic Kernel + AutoGen.",
        "Two core capabilities: Agents and Workflows.",
        "Code-first and fully programmable.",
        "Provider-agnostic — one interface, many backends.",
        "Python and C#.",
    ], size=14, gap=8)
    providers = ["Azure AI Foundry", "OpenAI", "Azure OpenAI", "Anthropic",
                 "Ollama", "AWS Bedrock", "Google Gemini"]
    py = 1.45
    add_textbox(s, 5.35, 1.05, 4.2, 0.3)
    _set_run(s.shapes[-1].text_frame.paragraphs[0].add_run(),
             "Swap with one env var:", 13, INK, bold=True)
    for i, pv in enumerate(providers):
        col = [INDIGO, MAGENTA, CYAN, PURPLE][i % 4]
        add_pill(s, 5.35 + (i % 2) * 2.25, py + (i // 2) * 0.55, 2.1, 0.42, pv, col, size=11)
    add_caption(s, 5.35, py + 4 * 0.55 + 0.05, 4.2,
                "MODEL_PROVIDER=…  → notebook code never changes", align=PP_ALIGN.LEFT)

    # 7) Section: Context Engineering
    section_slide(prs, "Context Engineering", "Deciding what the model sees")

    # 8) Context engineering
    text_and_diagram(
        prs, "Context Engineering",
        [
            "The model only sees what fits in its context window.",
            "Engineering that window is where real quality comes from.",
            ("Sessions — carry history across turns", 1),
            ("Context providers — inject facts before each run", 1),
            ("Memory — persist what matters", 1),
            ("Compaction — summarize so you never overflow", 1),
        ],
        "context-engineering.png",
        caption="What goes into the window each turn",
    )

    # 9) Section: The Agent Harness
    section_slide(prs, "The Agent Harness  ★", "Everything, assembled")

    # 10) The harness + diagram
    text_and_diagram(
        prs, "What is Agent Harness ?",
        [
            "Re-assembling the same machinery every time? The harness packages it.",
            "One factory call wires up:",
            ("Tool loop · history + persistence · compaction", 1),
            ("Todos (planning) · plan/execute modes", 1),
            ("Durable memory · skills · OpenTelemetry", 1),
            "Every battery can be disabled or replaced.",
        ],
        "agent-harness.png",
        caption="The 8 batteries of the harness",
    )

    # 11) Tools & function calling
    s = content_slide(prs, "Tools & Function Calling")
    add_bullets(s, 0.55, 1.35, 8.9, 2.3, [
        "A tool is a typed Python function + a docstring + the @tool decorator.",
        "The model decides when to call it; the framework runs it and loops.",
        "Approval gates put a human in the loop for risky actions:",
        ("approval_mode=\"always_require\" pauses for confirmation", 1),
        ("approval_mode=\"never_require\" runs automatically (demos)", 1),
    ], size=14, gap=7)
    for i, (lbl, col) in enumerate([("@tool", INDIGO), ("model calls it", MAGENTA),
                                    ("run + feed back", CYAN), ("approve?", PURPLE)]):
        add_pill(s, 0.7 + i * 2.1, 4.05, 1.85, 0.62, lbl, col, size=12)
    add_caption(s, 0.55, 4.85, 8.9,
                "Tools are the agent's hands — small, well-described, single-purpose.",
                align=PP_ALIGN.LEFT)

    # 12) Section: Orchestration
    section_slide(prs, "Multi-Agent Orchestration", "When one agent isn't enough")

    # 13) Orchestration patterns
    diagram_hero(
        prs, "Five Orchestration Patterns",
        "orchestration-patterns.png",
        bullets_below=[
            "Sequential · Concurrent · Handoff · Group Chat · Magentic — start with the simplest that fits.",
        ],
        caption="Specialized agents collaborating",
    )

    # 14) Section: Optimize & Operate
    section_slide(prs, "Optimize & Operate", "From demo to production")

    # 15) Evaluation
    text_and_diagram(
        prs, "Evaluating & Optimizing Agents",
        [
            "A demo that works once isn't a product.",
            "Measure quality so you can improve it on purpose:",
            ("Define checks (keyword, custom, model-graded)", 1),
            ("Run them over a query set", 1),
            ("Inspect failures → improve → re-run", 1),
            "raise_for_status() makes a regression break CI.",
        ],
        "evaluation-loop.png",
        caption="The evaluation loop",
    )

    # 16) Observability & ops
    text_and_diagram(
        prs, "Operationalizing — See Inside the Agent",
        [
            "Agents are non-deterministic and call external tools.",
            "Middleware = a control point around every model call:",
            ("Usage tracking (tokens / cost)", 1),
            ("Guardrails (redaction, injection checks)", 1),
            "OpenTelemetry gives end-to-end traces to any backend.",
        ],
        "observability.png",
        caption="Tracing, usage & guardrails",
    )

    # 17) Hands-on / resources
    s = content_slide(prs, "Hands-On — Build It Yourself")
    add_bullets(s, 0.55, 1.35, 8.9, 1.7, [
        "8 self-contained lab notebooks — M1 through M8.",
        "Provider-agnostic: pick a backend, the code stays the same.",
        "Runnable without an instructor — revisit anytime.",
    ], size=15, gap=8)
    add_rounded(s, 0.7, 3.25, 8.6, 1.0, CARD_BG)
    bx = add_textbox(s, 0.95, 3.42, 8.1, 0.7, anchor=MSO_ANCHOR.MIDDLE)
    p = bx.text_frame.paragraphs[0]
    _set_run(p.add_run(), "Live workshop site:  ", 15, INK, bold=True)
    _set_run(p.add_run(), "monuminu.github.io/Cypher-workshop26", 15, INDIGO, bold=True)
    p2 = bx.text_frame.add_paragraph()
    _set_run(p2.add_run(), "github.com/microsoft/agent-framework  ·  learn.microsoft.com/agent-framework",
             12, GRAY)

    # 18) Closing
    s = new_slide(prs, dark=True)
    box = add_textbox(s, 0.6, 2.0, 8.8, 1.0)
    p = box.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    _set_run(p.add_run(), "Thank You", 38, WHITE, bold=True)
    add_caption(s, 0.6, 3.15, 8.8, "Cypher 2026", size=18, color=WHITE)

    for index, slide in enumerate(prs.slides, 1):
        if index not in (1, len(prs.slides)):
            add_caption(slide, 0.55, 5.25, 8.9, f"Cypher 2026  ·  {index:02d}",
                        size=9, color=GRAY, align=PP_ALIGN.RIGHT)
    props = prs.core_properties
    props.title = "Cypher 2026 — Agentic AI"
    props.subject = "Microsoft Agent Framework workshop"
    props.author = "Manoranjan Rajguru"
    props.last_modified_by = "Manoranjan Rajguru"
    props.comments = ""
    props.keywords = "Cypher 2026, Agentic AI, Microsoft Agent Framework"
    prs.save(OUT)
    print(f"saved {OUT} with {len(prs.slides)} slides")


if __name__ == "__main__":
    build()
