"""Synthetic agent-written Python for offline control-flow tests only."""
import json


def workbook_code(*ids, completed=True):
    return '''import json
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
data = json.loads(Path('input.json').read_text(encoding='utf-8'))
contract = data['workbook_contract']
lookup = {s['id']: s for s in data['sessions']}
selected = ''' + repr(list(ids)) + '''
wb = Workbook()
wb.remove(wb.active)
for name in contract['sheets_in_order'][:2]:
    ws = wb.create_sheet(name)
    ws.append(contract['agenda_and_alternative_headers'])
    if name == 'My Agenda':
        for sid in selected:
            s = lookup[sid]
            ws.append([sid, s['start'][:10], s['start'][11:16], s['end'][11:16],
                       s['title'], s['hall'], s['speakers'], 'Synthetic reason', s['access'], s['source_url']])
ws = wb.create_sheet('Profile')
ws.append(['Setting','Value'])
for k,v in data['profile'].items():
    ws.append([k,json.dumps(v,ensure_ascii=False)])
ws = wb.create_sheet('Checks & Sources')
ws.append(['Check','Result'])
for k,v in contract['checks']['first_seven_rows']:
    ws.append([k, ''' + repr("Completed" if completed else "Draft — unresolved checks") + ''' if k == 'Status' else v])
for ws in wb:
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    for c in ws[1]:
        c.fill = PatternFill('solid',fgColor='3730A3')
        c.font = Font(name='Arial',color='FFFFFF',bold=True)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.data_type = 's'
            c.alignment = Alignment(wrap_text=True)
wb.save(contract['filename'])
wb.close()
check = load_workbook(contract['filename'])
print(check.sheetnames)
check.close()
'''


def python_step(*ids, completed=True):
    return ("python_execute", {"code": workbook_code(*ids, completed=completed)})


LOAD_SKILL = ("load_skill", {"skill_name": "xlsx"})
