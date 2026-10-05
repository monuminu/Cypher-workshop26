"""Synthetic fixtures only: no conference facts, network requests or paid models."""
import asyncio
import csv
import io
import json
import tempfile
import unittest
from contextlib import redirect_stdout
from dataclasses import asdict, replace
from datetime import date, datetime, timedelta
from pathlib import Path
from unittest.mock import patch
from urllib.parse import urlencode

from workshop_utils.agenda_source import (
    IST, SOURCE_URL, MANUAL_COLUMNS, Schedule, Session, fetch_live_schedule,
    import_reviewed_schedule, next_day, prepare_manual_import, session_from_calendar, parse_pdf_layout,
)
from workshop_utils.agenda import AgendaTools, Profile, validate, write_workbook, verify_workbook


def fixture():
    def s(sid, start, end, hall="Hall 1", **kwargs):
        return Session(sid, f"Synthetic {sid}", f"2026-10-08T{start}:00+05:30",
                       f"2026-10-08T{end}:00+05:30", hall, **kwargs)
    return Schedule([
        s("a", "10:00", "11:00"), s("b", "10:30", "10:50"),
        s("c", "11:00", "11:30", "Hall 2"), s("d", "11:05", "12:00", "Hall 2"),
        s("e", "14:00", "14:30", access="learning_or_vip"), s("lunch", "13:00", "13:30"),
    ], datetime.now(IST).isoformat(), "synthetic test", reviewed=True)


def proposal(*ids):
    return {"selections": [{"session_id": sid, "reason": "Synthetic teaching reason"} for sid in ids],
            "alternatives": []}


class AgendaTests(unittest.TestCase):
    def setUp(self):
        self.schedule, self.profile = fixture(), Profile("2026-10-08")
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def test_good_plan_and_buffer_boundary(self):
        self.assertTrue(validate(self.schedule, self.profile, proposal("a", "d")).ok)

    def test_overlap_and_hall_change(self):
        self.assertIn("overlaps", validate(self.schedule, self.profile, proposal("a", "b")).summary()["counts"])
        self.assertIn("hall_buffer", validate(self.schedule, self.profile, proposal("a", "c")).summary()["counts"])

    def test_long_interval_checks_all_overlaps(self):
        long = replace(self.schedule.sessions[0], end="2026-10-08T12:30:00+05:30")
        self.schedule.sessions[0] = long
        self.assertEqual(validate(self.schedule, self.profile, proposal("a", "b", "d")).summary()["counts"]["overlaps"], 2)

    def test_pass_lunch_revision_and_wrong_day(self):
        standard = replace(self.profile, pass_type="standard")
        self.assertIn("pass", validate(self.schedule, standard, proposal("e")).summary()["counts"])
        self.assertFalse(validate(self.schedule, self.profile, proposal("lunch")).ok)
        self.assertTrue(validate(self.schedule, self.profile, proposal("e")).ok)
        self.assertFalse(validate(self.schedule, self.profile.revised(), proposal("e")).ok)
        self.assertFalse(validate(self.schedule, replace(self.profile, day="2026-10-09"), proposal("a")).ok)

    def test_infeasible_mandatory_never_silently_relaxed(self):
        profile = replace(self.profile, must_attend=("a", "b"))
        self.assertFalse(validate(self.schedule, profile, proposal("a", "b")).ok)
        self.assertIn("mandatory", validate(self.schedule, profile, proposal("a")).summary()["counts"])

    def test_unknown_duplicate_empty_and_invalid_structure(self):
        for value in (proposal("missing"), proposal("a", "a"), proposal(), {"selections": "bad"}):
            self.assertFalse(validate(self.schedule, self.profile, value).ok)

    def test_missing_and_contradictory_source(self):
        self.schedule.sessions[0] = replace(self.schedule.sessions[0], hall="")
        self.assertFalse(validate(self.schedule, self.profile, proposal("a")).ok)
        self.assertTrue(replace(self.schedule.sessions[1], end="bad").issues())
        self.assertTrue(replace(self.schedule.sessions[1], access="unknown").issues())
        self.assertTrue(replace(self.schedule.sessions[1], start="2026-10-08T10:00:00").issues())

    def test_next_day_requires_explicit_past_selection(self):
        self.assertEqual(next_day(self.schedule, today=date(2026, 10, 7)), "2026-10-08")
        with self.assertRaises(ValueError):
            next_day(self.schedule, today=date(2026, 10, 8))
        self.assertEqual(next_day(self.schedule, "2026-10-08", today=date(2026, 10, 9)), "2026-10-08")

    def test_calendar_fields_and_timezone(self):
        href = "https://calendar.google.com/calendar/render?" + urlencode({
            "dates": "20261008T043000Z/20261008T053000Z", "text": "Synthetic agents",
            "location": "Hall 2 — Cypher 2026, Bengaluru, India",
            "details": "Category: Workshop\n\nSpeakers: Example Speaker\n\nA synthetic description.",
        })
        session = session_from_calendar("test", href, "Learning & VIP only")
        self.assertEqual(session.start, "2026-10-08T10:00:00+05:30")
        self.assertEqual(session.hall, "Hall 2")
        self.assertEqual(session.speakers, "Example Speaker")
        self.assertEqual(session.description, "A synthetic description.")
        self.assertEqual(session.access, "learning_or_vip")

    def test_manual_import_requires_fresh_reviewed_complete_rows(self):
        table = self.root / "sessions.tsv"
        with table.open("w", encoding="utf8", newline="") as f:
            writer = csv.DictWriter(f, fieldnames=MANUAL_COLUMNS, delimiter="\t")
            writer.writeheader()
            writer.writerows(asdict(s) for s in self.schedule.sessions)
        now = datetime.now(IST)
        imported = import_reviewed_schedule(table, captured_at=now.isoformat(), reviewed=True, now=now)
        self.assertEqual(imported.sessions, self.schedule.sessions)
        for reviewed, captured in ((False, now), (True, now - timedelta(hours=13)), (True, now + timedelta(minutes=1))):
            with self.assertRaises(ValueError):
                import_reviewed_schedule(table, captured_at=captured.isoformat(), reviewed=reviewed, now=now)
        raw, copied = prepare_manual_import(table, self.root / "review")
        self.assertIn("Synthetic a", raw.read_text(encoding="utf8"))
        self.assertEqual(copied.read_text(encoding="utf8"), table.read_text(encoding="utf8"))

    def test_copied_text_does_not_guess_grid_halls(self):
        source = self.root / "copy.txt"
        source.write_text("Hall 1  Hall 2\n10:00 Synthetic A  Synthetic B", encoding="utf8")
        _, table = prepare_manual_import(source, self.root / "review")
        self.assertEqual(len(table.read_text().splitlines()), 1)
        with self.assertRaises(ValueError):
            import_reviewed_schedule(table, captured_at=datetime.now(IST).isoformat(), reviewed=True)

    def test_pdf_extraction_and_empty_pdf(self):
        from pypdf import PdfWriter
        source = self.root / "source.pdf"
        writer = PdfWriter()
        writer.add_blank_page(width=400, height=400)
        writer.write(source)
        with self.assertRaises(ValueError):
            prepare_manual_import(source, self.root / "review")
        with patch("pypdf._page.PageObject.extract_text", return_value="Synthetic reviewed source text"):
            raw, _ = prepare_manual_import(source, self.root / "review")
            self.assertIn("Synthetic", raw.read_text())

    def test_labeled_pdf_layout_extracts_candidates_for_review(self):
        text = "Thursday, October 8, 2026\n"
        text += f"{'Time':<12}{'Hall':<12}{'Session':<50}Speakers\n"
        text += f"{'10:00':<12}{'Hall 2':<12}{'Synthetic title':<50}Example speaker\n"
        text += f"{'11:00':<12}{'':<12}{'WORKSHOP':<50}\n"
        text += f"{'':<12}{'':<12}{'Synthetic description':<50}\n"
        row = parse_pdf_layout(text)[0]
        self.assertEqual(row["title"], "Synthetic title")
        self.assertEqual(row["hall"], "Hall 2")
        self.assertEqual(row["end"], "2026-10-08T11:00:00+05:30")
        self.assertEqual(row["access"], "REVIEW_REQUIRED")

    def test_workbook_round_trip_and_tampered_facts(self):
        from openpyxl import load_workbook
        path = self.root / "agenda.xlsx"
        self.assertTrue(write_workbook(path, self.schedule, self.profile, proposal("a", "d")).ok)
        wb = load_workbook(path)
        self.assertEqual(wb["My Agenda"].freeze_panes, "A2")
        wb["My Agenda"]["F2"] = "Invented Hall"
        wb.save(path)
        wb.close()
        self.assertIn("workbook", verify_workbook(path, self.schedule, self.profile).summary()["counts"])

    def test_draft_export_and_stale_revision(self):
        tools = AgendaTools(self.schedule, self.profile, self.root)
        output = json.loads(tools.export_agenda(json.dumps(proposal("a", "b"))))
        self.assertFalse(output["valid"])
        tools.revise(self.profile.revised())
        self.assertFalse(json.loads(tools.inspect_workbook())["valid"])
        self.assertFalse(verify_workbook(self.root / "agenda-r1.xlsx", self.schedule, self.profile.revised()).ok)

    def test_excel_formula_injection_stays_text(self):
        from openpyxl import load_workbook
        value = proposal("a")
        value["selections"][0]["reason"] = '=HYPERLINK("https://example.invalid")'
        path = self.root / "agenda.xlsx"
        self.assertTrue(write_workbook(path, self.schedule, self.profile, value).ok)
        wb = load_workbook(path)
        self.assertEqual(wb["My Agenda"]["H2"].data_type, "s")
        wb.close()

    def test_profile_restart_roundtrip(self):
        path = self.root / "profile.json"
        self.profile.revised().save(path)
        self.assertEqual(Profile.load(path), self.profile.revised())


class SourceFailureTests(unittest.IsolatedAsyncioTestCase):
    async def test_live_failure_does_not_fallback(self):
        with patch("workshop_utils.agenda_source._fetch_browser", side_effect=RuntimeError("403")):
            with self.assertRaisesRegex(RuntimeError, "no cached"):
                await fetch_live_schedule()


if __name__ == "__main__":
    unittest.main()
